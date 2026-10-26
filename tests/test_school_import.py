from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
import pytest

from timelabs.domain import RuleError
from timelabs.school_import import read_school_workbook, read_rooms, source_report


@pytest.fixture
def school_files(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Группы — записи'
    sheet.append(['День', 'Урок', 'Класс', 'Профиль / слот', 'Предмет', 'Тип занятия', 'Цвет', 'Учитель'])
    rows = [
        ('Понедельник', 1, '5а', 'верхний слот', 'м', 'общее занятие', '#EECCAA', None),
        ('Понедельник', 1, '5а', 'нижний слот', 'м', 'общее занятие', '#EECCAA', None),
        ('Понедельник', 2, '5а', 'верхний слот', 'а/я · гр. 1', 'раздельно', '#AADDFF', None),
        ('Понедельник', 2, '5а', 'нижний слот', 'а/я · гр. 2', 'раздельно', '#FFAADD', None),
        ('Понедельник', 1, '11', 'соцгум', 'б', 'расшифровка 3 профилей, сверить', '#CCCCCC', None),
        ('Понедельник', 1, '11', 'химбио', 'б', 'расшифровка 3 профилей, сверить', '#CCCCCC', None),
        ('Понедельник', 1, '11', 'физмат', 'б', 'расшифровка 3 профилей, сверить', '#CCCCCC', None),
        ('Вторник', 3, '11', 'соцгум', 'ф', '11 класс: средний профиль предположительно с верхним — подтвердить', '#AAEEFF', None),
        ('Вторник', 3, '11', 'химбио', 'ф', '11 класс: средний профиль предположительно с верхним — подтвердить', '#AAEEFF', None),
        ('Вторник', 3, '11', 'физмат', 'А', 'отдельный профиль физмат', '#DDEEFF', None),
        ('Вторник', 6, '11', 'соцгум', 'ФК', 'общее занятие', '#DDFFFF', None),
        ('Вторник', 6, '11', 'химбио', 'ФК', 'общее занятие', '#DDFFFF', None),
        ('Вторник', 6, '11', 'физмат', 'ФК', 'общее занятие', '#DDFFFF', None),
    ]
    for row in rows:
        sheet.append(row)
    q = book.create_sheet('Вопросы по подгруппам')
    q.append(['Класс', 'День', 'Урок', 'Интерпретация', 'Вопрос', 'Ответ'])
    q.append(['11 класс', 'Понедельник', 1, 'б/б/б', 'Один урок или разные?', None])
    q.append(['11 класс', 'Вторник', 3, 'ф/А', 'Первые два профиля вместе?', 'да'])
    teachers = book.create_sheet('Учителя — снимок')
    teachers.append(['Заголовок'])
    teachers.append(['Учителя'])
    teachers.append(['Дни'])
    teachers.append(['Тестова А. А.', '5а', '5а'])
    book.create_sheet('Сокращения').append(['Код', 'Название'])
    book['Сокращения'].append(['м', 'Математика'])
    path = tmp_path / 'school.xlsx'
    book.save(path)
    rooms = tmp_path / 'rooms.docx'
    xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
           '<w:p><w:r><w:t>5а — Тестова (12) Другой (13)</w:t></w:r></w:p>'
           '<w:p><w:r><w:t>5а — Тестова (12) Другой (13)</w:t></w:r></w:p>'
           '</w:body></w:document>')
    with ZipFile(rooms, 'w') as z:
        z.writestr('word/document.xml', xml)
    return path, rooms


def test_common_event_is_one_identity_and_parallel_kept(school_files):
    source = read_school_workbook(*school_files)
    common = [e for e in source['events'] if e['school_class'] == '5а' and e['slot'] == 1]
    assert len(common) == 1
    assert common[0]['lane_runs'] == [[0, 1]]
    assert common[0]['profiles'] == ['верхний слот', 'нижний слот']
    assert len(common[0]['source_rows']) == 2
    parallel = [e for e in source['events'] if e['school_class'] == '5а' and e['slot'] == 2]
    assert len(parallel) == 2
    assert parallel[0]['subject'] == parallel[1]['subject'] == 'а/я'
    assert parallel[0]['lane_runs'] != parallel[1]['lane_runs']
    assert source['summary']['source_rows'] == sum(len(e['source_rows']) for e in source['events'])


def test_confirmed_profile_union_and_unanswered_not_merged(school_files):
    source = read_school_workbook(*school_files)
    unresolved = [e for e in source['events'] if e['school_class'] == '11' and e['weekday'] == 1]
    assert len(unresolved) == 3
    assert all('Один урок или разные?' in e['review'] for e in unresolved)
    confirmed = [e for e in source['events'] if e['school_class'] == '11' and e['weekday'] == 2 and e['slot'] == 3]
    assert len(confirmed) == 2
    assert next(e for e in confirmed if e['subject'] == 'ф')['lane_runs'] == [[0, 1]]
    full = [e for e in source['events'] if e['school_class'] == '11' and e['slot'] == 6]
    assert len(full) == 1
    assert full[0]['lane_runs'] == [[0, 1, 2]]


def test_candidates_are_not_confirmed_by_color_or_unique_snapshot(school_files):
    source = read_school_workbook(*school_files)
    event = next(e for e in source['events'] if e['school_class'] == '5а')
    assert event['teacher_candidates'] == ['Тестова А. А.']
    assert event['teacher'] is None
    assert event['room'] is None
    assert source['summary']['confirmed_teacher_events'] == 0
    assert source['rooms']['rooms'] == ['12', '13']
    assert len(source['rooms']['assignments']) == 2
    assert source['rooms']['assignments'][0]['status'] == 'proposed'
    assert source['rooms']['assignments'][0]['teacher'] is None
    assert 'НЕ ПОДТВЕРЖДЕНО' in source_report(source)


def test_deterministic_source_ids(school_files):
    first = read_school_workbook(*school_files)
    second = read_school_workbook(*school_files)
    assert first == second
    assert len({e['id'] for e in first['events']}) == len(first['events'])


def test_week_overview_keeps_same_columns_for_two_subgroups(school_files):
    from PySide6.QtWidgets import QApplication
    from timelabs.cards import CARDS_ROLE
    from timelabs.overview import WeekTable
    app = QApplication.instance() or QApplication([])
    source = read_school_workbook(*school_files)
    table = WeekTable(source)
    table.resize(1000, 400)
    table.show()
    app.processEvents()
    assert table.columnCount() == 45
    row5 = source['classes'].index('5а')
    common = table.item(row5, 0).data(CARDS_ROLE)
    separate = table.item(row5, 1).data(CARDS_ROLE)
    assert len(common) == 1 and common[0]['lane_span'] == 2
    assert len(separate) == 2
    assert table.columnWidth(0) == table.columnWidth(1)
    assert [c['lane_start'] for c in separate] == [0, 1]
    row11 = source['classes'].index('11')
    joined = table.item(row11, 9 + 2).data(CARDS_ROLE)
    assert sorted(c['lane_span'] for c in joined) == [1, 2]
    assert table.rowHeight(row11) > table.rowHeight(row5)
    assert table.item(row5, 0).toolTip().startswith('Математика')
    table.close()


def test_saved_review_does_not_replace_or_approve_active_schedule(school_files, tmp_path):
    from timelabs.storage import Store
    from timelabs.domain import demo_project
    source = read_school_workbook(*school_files)
    store = Store(tmp_path / 'review.db')
    store.create_user('admin', 'secure-review-password', 'admin')
    store.login('admin', 'secure-review-password')
    project = demo_project()
    lessons_before = project['lessons'].copy()
    project['source_workbook'] = source
    store.save(project, 0, 'school.import_review', 'catalog')
    version, loaded = store.load()
    assert loaded['lessons'] == lessons_before
    assert loaded['source_workbook']['status'] == 'review'
    assert not loaded['days']
    archive = tmp_path / 'school.tlabs'
    store.export_project(archive)
    store.import_project(archive)
    assert store.load()[1]['source_workbook'] == source
    store.db.close()


def test_negative_answer_does_not_confirm_assumed_union(school_files):
    from openpyxl import load_workbook
    path, rooms = school_files
    book = load_workbook(path)
    book['Вопросы по подгруппам']['F3'] = 'нет'
    book.save(path)
    book.close()
    source = read_school_workbook(path, rooms)
    events = [e for e in source['events'] if e['school_class'] == '11' and e['weekday'] == 2 and e['slot'] == 3]
    assert len(events) == 3
    assert source['summary']['unresolved_questions'] == 2
    assert source['records'][0]['source_values'][4] == 'м'


def test_explicit_nonadjacent_profiles_share_one_event(school_files):
    from timelabs.school_import import apply_clarification
    source = read_school_workbook(*school_files)
    source['teachers'].append('Примеров Б. Б.')
    clarification = dict(school_class='11', weekday=1, slot=1, source='Явный ответ пользователя',
                         partitions=[dict(profiles=['соцгум', 'физмат'], teacher='Тестова А. А.'),
                                     dict(profiles=['химбио'], teacher='Примеров Б. Б.')])
    corrected = apply_clarification(source, clarification)
    events = [e for e in corrected['events'] if e['school_class'] == '11' and e['weekday'] == 1]
    assert len(events) == 2
    shared = next(e for e in events if e['teacher'] == 'Тестова А. А.')
    assert shared['lanes'] == [0, 2]
    assert shared['lane_runs'] == [[0], [2]]
    assert len(shared['source_rows']) == 2
    assert corrected['summary']['unresolved_questions'] == 0
    assert source['summary']['unresolved_questions'] == 1
    assert corrected['class_profiles']['11'] == ['соцгум', 'химбио', 'физмат']
    assert source['events'] != corrected['events']
    assert sum(len(e['source_rows']) for e in corrected['events']) == len(source['records'])


def test_manual_clarification_rejects_missing_profiles_and_unknown_teacher(school_files):
    from timelabs.school_import import apply_clarification
    source = read_school_workbook(*school_files)
    with pytest.raises(RuleError, match='ровно один'):
        apply_clarification(source, dict(school_class='11', weekday=1, slot=1,
            partitions=[dict(profiles=['соцгум'], teacher='Тестова А. А.')]))
    with pytest.raises(RuleError, match='справочника'):
        apply_clarification(source, dict(school_class='11', weekday=1, slot=1,
            partitions=[dict(profiles=['соцгум','химбио','физмат'], teacher='Несуществующий')]))


def test_profile_class_declarations_and_single_fizmat_class(school_files):
    from openpyxl import load_workbook
    path, rooms = school_files
    book = load_workbook(path)
    book['Группы — записи'].append(['Понедельник', 1, '10а', 'соцгум', 'м', 'общее занятие', '#FFFFFF', None])
    book['Группы — записи'].append(['Понедельник', 1, '10а', 'химбио', 'м', 'общее занятие', '#FFFFFF', None])
    book['Группы — записи'].append(['Понедельник', 1, '10б', 'физмат', 'м', 'весь класс', '#FFFFFF', None])
    book.save(path)
    book.close()
    source = read_school_workbook(path, rooms)
    assert source['class_profiles'] == {'11':['соцгум','химбио','физмат'], '10а':['соцгум','химбио'], '10б':['физмат']}
    assert source['class_lanes']['10а'] == 2
    assert source['class_lanes']['10б'] == 1
    assert len([e for e in source['events'] if e['school_class']=='10б']) == 1


def test_review_ui_can_confirm_outer_profiles_without_mutating_cancelled_input(school_files):
    from PySide6.QtCore import QTimer, Qt
    from PySide6.QtWidgets import QApplication, QDialog, QDialogButtonBox, QTableWidget
    from timelabs.cards import CARDS_ROLE
    from timelabs.overview import SourceDialog
    app = QApplication.instance() or QApplication([])
    source = read_school_workbook(*school_files)
    source['teachers'].append('Примеров Б. Б.')
    dialog = SourceDialog(source, can_save=True)
    row = source['classes'].index('11')
    def fill():
        editor = app.activeModalWidget()
        table = editor.findChild(QTableWidget)
        for i, group in enumerate([1,2,1]):
            table.cellWidget(i,1).setValue(group)
            teacher = table.cellWidget(i,2)
            teacher.setCurrentIndex(teacher.findData('Примеров Б. Б.' if i == 1 else 'Тестова А. А.'))
        buttons = editor.findChild(QDialogButtonBox)
        buttons.button(QDialogButtonBox.StandardButton.Save).click()
    QTimer.singleShot(100, fill)
    detail_dialog = QDialog(dialog)
    dialog.edit_slot(row, 0, detail_dialog)
    assert dialog.source['summary']['unresolved_questions'] == 0
    assert source['summary']['unresolved_questions'] == 1
    cards = dialog.table.item(row,0).data(CARDS_ROLE)
    assert len(cards) == 3
    assert cards[0]['id'] == cards[1]['id']
    assert len(set(c['id'] for c in cards)) == 2
    dialog.reject()
    dialog.deleteLater()
    app.processEvents()
