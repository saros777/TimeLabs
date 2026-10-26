from html import escape

from openpyxl import Workbook, load_workbook

from .domain import RuleError, uid, validate

HEADERS = ['День', 'Урок', 'Предмет ID', 'Группы ID', 'Учителя ID', 'Кабинет ID', 'Длительность', 'Закреплён', 'Комментарий']


def export_xlsx(project, lessons, path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Расписание'
    sheet.append(HEADERS)
    for l in lessons:
        sheet.append([l['weekday'], l['slot'], l['subject'], ','.join(l['groups']), ','.join(l['teachers']),
                      l['room'], l.get('duration', 1), int(l.get('locked', False)), l.get('comment', '')])
    sheet.freeze_panes = 'A2'
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 24
    for kind in ('teachers', 'groups', 'rooms', 'subjects'):
        ws = book.create_sheet(kind)
        ws.append(['ID', 'Название'])
        for key, value in project[kind].items():
            ws.append([key, value['name']])
    book.save(path)


def preview_xlsx(project, path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book['Расписание'] if 'Расписание' in book.sheetnames else book.active
        rows = iter(sheet.values)
        headers = list(next(rows))[:len(HEADERS)]
        if headers != HEADERS:
            raise RuleError('Нужен согласованный формат XLSX. Сначала экспортируйте шаблон; произвольная школьная книга требует отдельного сопоставления.')
        lessons, issues = [], []
        for i, row in enumerate(rows, 2):
            if not any(x is not None for x in row):
                continue
            try:
                day, slot, subject, groups, teachers, room, duration, locked, comment = (list(row) + [None] * 9)[:9]
                if locked not in (0, 1, False, True, None):
                    raise RuleError('Закреплён: только 0 или 1')
                lessons.append(dict(id=uid(), weekday=int(day), slot=int(slot), subject=str(subject),
                                    groups=[x.strip() for x in str(groups).split(',')],
                                    teachers=[x.strip() for x in str(teachers).split(',')], room=str(room),
                                    duration=int(duration or 1), locked=bool(locked), comment=str(comment or ''), status='planned'))
            except (ValueError, TypeError, RuleError) as exc:
                issues.append(f'Строка {i}: {exc}')
        issues.extend(validate(project, lessons))
        return lessons, issues
    finally:
        book.close()


def schedule_html(project, lessons, title, subtitle=''):
    rows = []
    for l in sorted(lessons, key=lambda x: (x['weekday'], x['slot'], x['groups'])):
        values = [str(l['slot']), project['subjects'][l['subject']]['name'],
                  ', '.join(project['groups'][g]['name'] for g in l['groups']),
                  ', '.join(project['teachers'][t]['name'] for t in l['teachers']),
                  project['rooms'][l['room']]['name'],
                  {'cancelled': 'ОТМЕНА', 'makeup': 'ВОСПОЛНЕНИЕ'}.get(l.get('status'), '') + ' ' + l.get('comment', '')]
        rows.append('<tr>' + ''.join(f'<td>{escape(v)}</td>' for v in values) + '</tr>')
    return ('<html><head><meta charset="utf-8"><style>body{font-family:DejaVu Sans,Arial;font-size:10pt}'
            'table{border-collapse:collapse;width:100%}td,th{border:1px solid #8796a8;padding:6px;text-align:left}'
            '</style></head><body>' + f'<h1>{escape(title)}</h1><p>{escape(subtitle)}</p>' +
            '<table><thead><tr><th>№</th><th>Предмет</th><th>Группы</th><th>Учителя</th><th>Кабинет</th><th>Изменения</th></tr></thead>' +
            '<tbody>' + ''.join(rows) + '</tbody></table></body></html>')
