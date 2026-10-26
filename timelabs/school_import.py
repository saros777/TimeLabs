"""Traceable draft reconciliation of the customer's school workbook.
No teacher, room, pupil count, qualification or capacity is invented.
"""
from collections import defaultdict
from hashlib import sha256
from pathlib import Path
import re
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook

from .domain import RuleError

DAYS = ['Понедельник', 'Вторник', 'Среда', 'Четверг', 'Пятница']
PROFILE_ORDER = {'весь класс': 0, 'верхний слот': 0, 'нижний слот': 1,
                 'соцгум': 0, 'химбио': 1, 'физмат': 2}
CLASS_PROFILES = {'11': ['соцгум', 'химбио', 'физмат'],
                  '10а': ['соцгум', 'химбио'], '10б': ['физмат']}


def profile_lane(school_class, profile):
    declared = CLASS_PROFILES.get(school_class, [])
    return declared.index(profile) if profile in declared else PROFILE_ORDER[profile]


def stable_id(*parts):
    return sha256('|'.join(map(str, parts)).encode()).hexdigest()[:24]


def name_key(text):
    return re.sub(r'[^а-яa-z]', '', text.casefold().replace('ё', 'е'))


def class_sort(text):
    match = re.fullmatch(r'(\d+)(.*)', text)
    return (int(match[1]), match[2]) if match else (100, text)


def answer_confirms(question):
    if not question or not question.get('answer'):
        return False
    if question.get('confirmed_manually'):
        return True
    answer = str(question['answer']).strip()
    return (answer.casefold() == 'да' or
            all(fragment in answer for fragment in ('сверху', 'нижн', 'Бакай', 'Широбоков')))


def read_rooms(path, teacher_names):
    """Read paragraphs once; Word drawing/textbox repetitions are deduplicated."""
    with ZipFile(path) as archive:
        if archive.getinfo('word/document.xml').file_size > 8_000_000:
            raise RuleError('Документ кабинетов слишком большой')
        root = ElementTree.fromstring(archive.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    paragraphs = list(dict.fromkeys(''.join(t.text or '' for t in p.findall('.//w:t', ns))
                                    for p in root.findall('.//w:p', ns)))
    assignments, rooms = [], set()
    for text in paragraphs:
        match = re.match(r'\s*(\d+[а-я]?)\s*[—–-]\s*(.*)', text)
        if not match:
            continue
        school_class, tail = match.groups()
        for person, room in re.findall(r'([^()]+?)\s*\((\d+)\s*(?:каб)?\s*\)', tail):
            person = person.strip()
            rooms.add(room)
            key = name_key(person)
            surname_only = len(person.split()) == 1
            exact = [n for n in teacher_names if name_key(n) == key]
            proposals = exact or [n for n in teacher_names if name_key(n.split()[0]) == key] if surname_only else exact
            # Surname-only matches remain proposals even when unique.
            status = 'exact' if exact else 'proposed' if proposals else 'unresolved'
            assignments.append(dict(school_class=school_class, source_name=person, room=room,
                                    teacher=exact[0] if len(exact) == 1 else None,
                                    candidates=proposals, status=status, preference_only=True))
    return dict(source=Path(path).name, sha256=sha256(Path(path).read_bytes()).hexdigest(),
                rooms=sorted(rooms, key=int), assignments=assignments,
                unknown_fields=['Вместимость', 'Оснащение', 'Доступность'])


def read_school_workbook(path, room_path=None):
    book = load_workbook(path, data_only=True, read_only=True)
    try:
        required = {'Группы — записи', 'Вопросы по подгруппам', 'Учителя — снимок', 'Сокращения'}
        if not required <= set(book.sheetnames):
            raise RuleError('Не обнаружены листы уточнённой книги расписания')
        abbreviations = {str(k): str(v) for k, v in list(book['Сокращения'].values)[1:] if k}
        questions = []
        for number, row in enumerate(list(book['Вопросы по подгруппам'].values)[1:], 2):
            if not row[0]:
                continue
            questions.append(dict(row=number, school_class=str(row[0]).replace(' класс', ''),
                                  day=row[1], slot=row[2], interpretation=row[3], question=row[4],
                                  answer=row[5] if len(row) > 5 else None))
        questions_by_slot = {(q['school_class'], q['day'], q['slot']): q for q in questions}
        teacher_names, occupancy = [], defaultdict(list)
        for row in list(book['Учителя — снимок'].values)[3:]:
            if not row[0]:
                continue
            teacher_names.append(row[0])
            for offset, value in enumerate(row[1:46]):
                if value:
                    occupancy[(DAYS[offset // 9], offset % 9 + 1, str(value).strip())].append(row[0])
        records, errors = [], []
        for row_number, values in enumerate(list(book['Группы — записи'].values)[1:], 2):
            day, slot, school_class, profile, raw_subject, kind, color, teacher = values[:8]
            if day is None:
                continue
            if day not in DAYS or not isinstance(slot, int) or not 1 <= slot <= 9 or profile not in PROFILE_ORDER:
                errors.append(f'Группы — записи!{row_number}: неизвестный день, номер урока или профиль')
                continue
            subject = str(raw_subject).split(' · ')[0].strip()
            school_class = str(school_class).strip()
            q = questions_by_slot.get((school_class, day, slot))
            review = []
            if q and not answer_confirms(q):
                review.append(q['question'])
            elif ('сверить' in str(kind) or 'подтвердить' in str(kind)) and not answer_confirms(q):
                review.append(str(kind))
            # Explicit answers in the workbook disambiguate the two history/social studies teachers.
            answer = str(q['answer']) if q and q['answer'] else ''
            teacher_source = 'Явная колонка преподавателя' if teacher else None
            if (school_class == '11' and subject in {'и', 'о'} and
                    'Цвета разные' in answer and 'Бакай' in answer and 'Широбоков' in answer):
                surname = ('Бакай' if profile == 'соцгум' else 'Широбокова') if subject == 'и' else ('Широбокова' if profile == 'соцгум' else 'Бакай')
                named = [n for n in teacher_names if (n.startswith('Бакай ') if surname == 'Бакай' else name_key(n) == name_key('Широбокова О. А.'))]
                if len(named) == 1:
                    teacher = named[0]
                    teacher_source = f'Ответ на листе вопросов, строка {q["row"]}'
            records.append(dict(row=row_number, source_values=list(values[:8]), weekday=DAYS.index(day) + 1, day=day, slot=slot,
                                school_class=school_class, profile=profile, subject=subject,
                                subject_name=abbreviations.get(subject, subject), kind=str(kind), color=color or '#E2EDFA',
                                teacher=teacher, teacher_source=teacher_source, review=review,
                                candidates=occupancy.get((day, slot, school_class), [])))
        by_slot = defaultdict(list)
        for row in records:
            by_slot[(row['school_class'], row['weekday'], row['slot'])].append(row)
        events = []
        class_lanes = defaultdict(int)
        for row in records:
            class_lanes[row['school_class']] = max(class_lanes[row['school_class']], profile_lane(row['school_class'], row['profile']) + 1)
        for (school_class, weekday, slot), rows in by_slot.items():
            partitions = defaultdict(list)
            for row in rows:
                kind = row['kind']
                question = questions_by_slot.get((school_class, row['day'], slot))
                confirmed_answer = answer_confirms(question)
                merge = ('общее занятие' in kind or 'объединение' in kind or
                         ('предположительно' in kind and confirmed_answer)) and not row['review']
                # Do not merge equal subject text when the source says separate or uncertain.
                merge_key = (row['subject'], row['teacher']) if merge else (row['subject'], row['teacher'], row['profile'])
                partitions[merge_key].append(row)
            for parts in partitions.values():
                first = parts[0]
                lanes = sorted({profile_lane(school_class, r['profile']) for r in parts})
                if lanes != list(range(lanes[0], lanes[-1] + 1)):
                    # Noncontiguous joined profiles retain one identity and multiple paint bands.
                    runs = [[i] for i in lanes]
                else:
                    runs = [lanes]
                reasons = sorted({reason for r in parts for reason in r['review']})
                if not first['teacher']:
                    reasons.append('Преподаватель не подтверждён')
                reasons.append('Фактический кабинет не назначен')
                events.append(dict(id=stable_id('school', school_class, weekday, slot, *(r['row'] for r in parts)),
                                   school_class=school_class, weekday=weekday, slot=slot, subject=first['subject'],
                                   subject_name=first['subject_name'], profiles=[r['profile'] for r in parts],
                                   lanes=lanes, lane_runs=runs, lane_total=class_lanes[school_class],
                                   teacher=first['teacher'], teacher_source=first['teacher_source'],
                                   teacher_candidates=first['candidates'], room=None, color=first['color'],
                                   source_rows=[r['row'] for r in parts], review=reasons,
                                   source_kinds=sorted({r['kind'] for r in parts})))
        result = dict(format=1, source=Path(path).name, sha256=sha256(Path(path).read_bytes()).hexdigest(),
                      sheets=list(book.sheetnames), records=records, events=events,
                      classes=sorted(class_lanes, key=class_sort), class_lanes=dict(class_lanes),
                      class_profiles={k: v for k, v in CLASS_PROFILES.items() if k in class_lanes},
                      teachers=teacher_names, abbreviations=abbreviations, questions=questions,
                      errors=errors, status='review', rooms=read_rooms(room_path, teacher_names) if room_path else None)
        result['summary'] = dict(source_rows=len(records), events=len(events), classes=len(class_lanes),
                                 teachers=len(teacher_names), confirmed_teacher_events=sum(bool(e['teacher']) for e in events),
                                 unresolved_questions=sum(not answer_confirms(q) for q in questions), errors=len(errors))
        return result
    finally:
        book.close()


def source_report(source):
    s = source['summary']
    lines = ['# Сверка школьной книги', '', f"Файл: {source['source']}",
             f"Исходных строк: {s['source_rows']}; единых событий: {s['events']}; классов: {s['classes']}; преподавателей: {s['teachers']}.",
             f"Явно подтверждён преподаватель у {s['confirmed_teacher_events']} событий. Кандидаты из снимка — только предложения.",
             'Кабинеты не назначены автоматически; вместимость и оснащение не заданы.',
             'Цвет сохранён для отображения и не подтверждает личность преподавателя.', '', '## Вопросы профилей']
    for school_class, profiles in source.get('class_profiles', {}).items():
        lines.insert(-1, f"Порядок профилей {school_class} сверху вниз: {', '.join(profiles)}" + (' (единый класс без деления).' if len(profiles) == 1 else '.'))
    for q in source['questions']:
        lines.append(f"- {q['school_class']} · {q['day']} · {q['slot']}: {q['question']} Ответ: {q['answer'] or 'НЕ ПОДТВЕРЖДЕНО'}")
    if source['rooms']:
        lines += ['', '## Закрепления кабинетов (предпочтения)']
        for a in source['rooms']['assignments']:
            lines.append(f"- {a['school_class']} · {a['source_name']} · {a['room']}: {a['status']}; варианты: {', '.join(a['candidates']) or 'нет точного совпадения'}")
    lines += ['', '## Ошибки структуры'] + (source['errors'] or ['Не обнаружены. Это не означает, что школьные данные полностью подтверждены.'])
    return '\n'.join(lines) + '\n'


def apply_clarification(source, clarification):
    """Apply an explicit human decision to the review copy, never the active plan.

    Each partition names the profiles sharing one event and its confirmed teacher.
    Nonadjacent profiles retain one event ID with multiple visual bands.
    """
    from copy import deepcopy
    result = deepcopy(source)
    school_class = clarification['school_class']
    weekday, slot = clarification['weekday'], clarification['slot']
    rows = [r for r in result['records'] if (r['school_class'], r['weekday'], r['slot']) == (school_class, weekday, slot)]
    by_profile = {r['profile']: r for r in rows}
    if not rows or len(by_profile) != len(rows):
        raise RuleError('Не найдены однозначные строки профилей для уточнения')
    supplied = [p for partition in clarification['partitions'] for p in partition['profiles']]
    if len(set(supplied)) != len(supplied) or set(supplied) != set(by_profile):
        raise RuleError('Каждый профиль слота должен быть указан ровно один раз')
    events = []
    for partition in clarification['partitions']:
        parts = sorted([by_profile[p] for p in partition['profiles']], key=lambda r: profile_lane(school_class, r['profile']))
        if len({r['subject'] for r in parts}) != 1:
            raise RuleError('Нельзя объединить в одно занятие разные предметы')
        teacher = partition['teacher']
        if teacher not in result['teachers']:
            raise RuleError('Выберите преподавателя из справочника, без приблизительного сопоставления')
        lanes = [profile_lane(school_class, r['profile']) for r in parts]
        runs = []
        for lane in lanes:
            if runs and lane == runs[-1][-1] + 1:
                runs[-1].append(lane)
            else:
                runs.append([lane])
        first = parts[0]
        events.append(dict(id=stable_id('school', school_class, weekday, slot, *(r['row'] for r in parts)),
                           school_class=school_class, weekday=weekday, slot=slot,
                           subject=first['subject'], subject_name=first['subject_name'],
                           profiles=[r['profile'] for r in parts], lanes=lanes, lane_runs=runs,
                           lane_total=result['class_lanes'][school_class], teacher=teacher,
                           teacher_source=clarification.get('source', 'Ручная сверка пользователя'),
                           teacher_candidates=first['candidates'], room=None, color=first['color'],
                           source_rows=[r['row'] for r in parts], review=['Фактический кабинет не назначен'],
                           source_kinds=sorted({r['kind'] for r in parts}), confirmed_manually=True))
    result['events'] = [e for e in result['events'] if (e['school_class'], e['weekday'], e['slot']) != (school_class, weekday, slot)] + events
    result.setdefault('clarifications', []).append(deepcopy(clarification))
    for q in result['questions']:
        if (q['school_class'], q['day'], q['slot']) == (school_class, DAYS[weekday - 1], slot):
            q.update(answer=clarification.get('answer') or '; '.join(', '.join(p['profiles']) + ' — ' + p['teacher'] for p in clarification['partitions']),
                     confirmed_manually=True, confirmation_source=clarification.get('source', 'Ручная сверка пользователя'))
    result['summary'].update(events=len(result['events']),
                             confirmed_teacher_events=sum(bool(e['teacher']) for e in result['events']),
                             unresolved_questions=sum(not answer_confirms(q) for q in result['questions']))
    return result
