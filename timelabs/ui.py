from __future__ import annotations

import json
import multiprocessing as mp
import os
import sys
import time
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path

from PySide6.QtCore import QDate, QEvent, QLockFile, QMimeData, Qt, QTimer, Signal, QSize
from PySide6.QtGui import QAction, QColor, QDrag, QFont, QKeySequence, QTextDocument, QPageLayout
from PySide6.QtWidgets import (
    QApplication, QCheckBox, QComboBox, QDateEdit, QDialog, QDialogButtonBox, QFileDialog,
    QFormLayout, QHBoxLayout, QHeaderView, QInputDialog, QLabel, QLineEdit, QListWidget,
    QMainWindow, QMessageBox, QPushButton, QSpinBox, QSplitter, QTableWidget, QListWidgetItem, QFrame, QStackedWidget, QMenu,
    QTableWidgetItem, QTextEdit, QVBoxLayout, QWidget, QProgressDialog, QToolTip,
)
from PySide6.QtPrintSupport import QPrinter, QPrintPreviewDialog

from .documents import export_xlsx, preview_xlsx, schedule_html
from .domain import RuleError, day_lessons, demo_project, uid, validate
from .services import Service
from .solver import solve_pair
from .storage import ROLES, Store
from .weekly import generate_week
from .catalog import CatalogDialog
from .cards import CARDS_ROLE, PREVIEW_ROLE, LessonDelegate, layout_cards
from .school_import import apply_clarification, read_school_workbook, class_sort
from .overview import SourceDialog
from .school_transfer import transfer_school
from .desk import DeskTable, STYLE

DAYS = ['Понедельник', 'Вторник', 'Среда', 'Четверг', 'Пятница', 'Суббота', 'Воскресенье']
STATUSES = {'draft': 'Черновик', 'checked': 'Проверено', 'approved': 'Утверждено', 'published': 'Опубликовано'}
ROLE_NAMES = {'admin': 'Администратор', 'planner': 'Составитель', 'dispatcher': 'Диспетчер', 'viewer': 'Просмотр'}


def alert(parent, error):
    QMessageBox.warning(parent, 'TimeLabs', str(error))


def confirm(parent, text):
    return QMessageBox.question(parent, 'Подтверждение', text) == QMessageBox.StandardButton.Yes


class Form(QDialog):
    def __init__(self, parent, title):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.setMinimumWidth(440)
        self.form = QFormLayout(self)
        self.buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        self.buttons.accepted.connect(self.accept)
        self.buttons.rejected.connect(self.reject)

    def finish(self):
        self.form.addRow(self.buttons)
        return self.exec() == QDialog.DialogCode.Accepted

    def text(self, label, value='', password=False):
        edit = QLineEdit(value)
        if password:
            edit.setEchoMode(QLineEdit.EchoMode.Password)
        self.form.addRow(label, edit)
        return edit

    def combo(self, label, entries):
        edit = QComboBox()
        for key, name in entries:
            edit.addItem(name, key)
        self.form.addRow(label, edit)
        return edit

    def number(self, label, value=1, maximum=12):
        edit = QSpinBox()
        edit.setRange(1, maximum)
        edit.setValue(value)
        self.form.addRow(label, edit)
        return edit

    def check(self, label, checked=False):
        edit = QCheckBox(label)
        edit.setChecked(checked)
        self.form.addRow(edit)
        return edit


class Login(Form):
    def __init__(self, store, parent=None):
        first = not store.initialized()
        super().__init__(parent, 'Первый запуск · создайте администратора' if first else 'Вход в TimeLabs')
        self.store = store
        self.first = first
        self.login_edit = self.text('Логин', 'admin' if first else '')
        self.password_edit = self.text('Пароль (минимум 10 символов)', password=True)
        self.repeat = self.text('Повтор пароля', password=True) if first else None
        self.form.addRow(QLabel('Локальная учётная запись · подключение к интернету не требуется'))
        self.form.addRow(self.buttons)

    def accept(self):
        try:
            if self.first:
                if self.password_edit.text() != self.repeat.text():
                    raise RuleError('Пароли не совпадают')
                self.store.create_user(self.login_edit.text(), self.password_edit.text(), 'admin')
                self.first = False
            self.store.login(self.login_edit.text(), self.password_edit.text())
            super().accept()
        except Exception as exc:
            alert(self, exc)


class Grid(QTableWidget):
    moved = Signal(str, int)

    def __init__(self, parent):
        super().__init__(parent)
        self.setDragEnabled(True)
        self.setAcceptDrops(True)
        self.setDragDropMode(QTableWidget.DragDropMode.DragDrop)
        self.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.setSelectionMode(QTableWidget.SelectionMode.SingleSelection)
        self.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.horizontalHeader().setMinimumSectionSize(130)
        self.setWordWrap(True)
        self.setItemDelegate(LessonDelegate(self))
        self.hit_lesson = None

    def lesson_at(self, position):
        item = self.itemAt(position)
        if not item:
            return None
        rect = self.visualItemRect(item)
        for card, band in layout_cards(item.data(CARDS_ROLE) or [], rect.width(), rect.height()):
            if band.contains(position.x() - rect.x(), position.y() - rect.y()):
                return card['id']
        return None

    def mousePressEvent(self, event):
        self.hit_lesson = self.lesson_at(event.position().toPoint())
        super().mousePressEvent(event)

    def mouseDoubleClickEvent(self, event):
        self.hit_lesson = self.lesson_at(event.position().toPoint())
        super().mouseDoubleClickEvent(event)

    def viewportEvent(self, event):
        if event.type() == QEvent.Type.ToolTip:
            item = self.itemAt(event.pos())
            if item and item.data(PREVIEW_ROLE) is None:
                key = self.lesson_at(event.pos())
                card = next((c for c in item.data(CARDS_ROLE) or [] if c['id'] == key), None)
                if card:
                    QToolTip.showText(event.globalPos(), card.get('tooltip', card['label']), self)
                    return True
        return super().viewportEvent(event)

    def startDrag(self, supported):
        item = self.currentItem()
        if not item or not item.data(Qt.ItemDataRole.UserRole):
            return
        ids = item.data(Qt.ItemDataRole.UserRole)
        selected = self.hit_lesson if self.hit_lesson in ids else ids[0] if len(ids) == 1 else None
        if selected is None:
            return
        mime = QMimeData()
        mime.setData('application/x-timelabs-lesson', selected.encode())
        self.window().preview_move(selected)
        drag = QDrag(self)
        drag.setMimeData(mime)
        drag.exec(Qt.DropAction.MoveAction)
        self.window().refresh()

    def dragEnterEvent(self, event):
        if event.mimeData().hasFormat('application/x-timelabs-lesson'):
            event.acceptProposedAction()

    def dragMoveEvent(self, event):
        event.acceptProposedAction()

    def dropEvent(self, event):
        row = self.rowAt(int(event.position().y()))
        if row >= 0 and event.mimeData().hasFormat('application/x-timelabs-lesson'):
            self.moved.emit(bytes(event.mimeData().data('application/x-timelabs-lesson')).decode(), row + 1)
            event.acceptProposedAction()


def _solve_worker(connection, project, day, weekly=False):
    try:
        connection.send(('ok', generate_week(project) if weekly else solve_pair(project, day)))
    except Exception as exc:
        connection.send(('error', str(exc)))
    finally:
        connection.close()


class Window(QMainWindow):
    def __init__(self, store):
        super().__init__()
        self.store, self.service = store, Service(store)
        self.setWindowTitle('TimeLabs · Школьное расписание')
        self.resize(1440, 900)
        self.undo_stack, self.redo_stack = [], []
        self.worker = None
        self.last_activity = time.monotonic()
        self.locking = False
        self.timer = QTimer(self)
        self.timer.timeout.connect(self.check_idle)
        self.timer.start(1000)
        QApplication.instance().installEventFilter(self)
        central = QWidget()
        layout = QVBoxLayout(central)
        layout.setContentsMargins(0,0,0,0)
        layout.setSpacing(0)
        nav=QFrame()
        nav.setObjectName('navigation')
        nav_bar=QHBoxLayout(nav)
        nav_bar.setContentsMargins(0,0,12,0)
        nav_bar.setSpacing(0)
        brand=QLabel('▦  TimeLabs')
        brand.setObjectName('brand')
        nav_bar.addWidget(brand)
        for label,callback in [('Расписание', lambda: self.layout_mode.setCurrentIndex(0)),
                               ('Замены', self.open_replacements), ('Контроль расписания', self.open_control),
                               ('Справочники', self.catalog), ('История изменений',self.history)]:
            button=QPushButton(label)
            if label=='Расписание':
                button.setCheckable(True)
                button.setChecked(True)
            button.clicked.connect(callback)
            nav_bar.addWidget(button)
        database=QPushButton('База и резервные копии')
        database_menu=QMenu(database)
        for label,callback in [('Импорт проекта…',self.import_archive),('Импорт Excel…',self.import_excel),
                               ('Экспорт проекта…',self.export),('Резервная копия…',self.backup),
                               ('Загрузить сохранённую книгу…',self.activate_school),
                               ('Извлечь назначения учителей…',self.extract_teacher_assignments),
                               ('Публикация дня',lambda: self.transition('published'))]:
            database_menu.addAction(label,callback)
        database.setMenu(database_menu)
        nav_bar.addWidget(database)
        nav_bar.addStretch()
        layout.addWidget(nav)
        content=QWidget()
        body=QVBoxLayout(content)
        body.setContentsMargins(10,10,10,8)
        body.setSpacing(9)
        project_bar=QHBoxLayout()
        self.banner=QLabel()
        self.banner.setObjectName('muted')
        project_bar.addWidget(self.banner,1)
        self.saved=QLabel('На диске · сохранено')
        self.saved.setStyleSheet('color:#3A7C59')
        project_bar.addWidget(self.saved)
        for label,callback in [('Импорт',self.import_excel),('Резервная копия',self.backup)]:
            button=QPushButton(label)
            button.clicked.connect(callback)
            project_bar.addWidget(button)
        body.addLayout(project_bar)
        tools=QHBoxLayout()
        heading=QLabel('Расписание')
        heading.setObjectName('heading')
        tools.addWidget(heading,1)
        for label,callback in [('＋ Занятие',self.add_lesson),('Отсутствие учителя',self.absence),
                               ('Пересчитать день',self.calculate),('Печать / PDF',self.pdf),
                               ('Проверить',lambda: self.transition('checked')),
                               ('Утвердить',lambda: self.transition('approved'))]:
            button=QPushButton(label)
            if label=='Пересчитать день':button.setObjectName('primary')
            button.clicked.connect(callback)
            tools.addWidget(button)
        body.addLayout(tools)
        caption=QLabel('Классы сверху вниз · дни и уроки слева направо · двойной щелчок — редактирование')
        caption.setObjectName('muted')
        body.addWidget(caption)
        bar=QHBoxLayout()
        self.layout_mode=QComboBox()
        self.layout_mode.addItems(['Неделя','День'])
        self.scope=QComboBox()
        self.scope.addItems(['На дату','Базовая неделя'])
        self.scope.setCurrentIndex(1)
        self.date_edit=QDateEdit(QDate.currentDate())
        self.date_edit.setCalendarPopup(True)
        self.date_edit.setDisplayFormat('dd.MM.yyyy')
        self.weekday=QComboBox()
        self.weekday.addItems(DAYS)
        self.weekday.setCurrentIndex(date.today().weekday())
        previous=self.previous_day=QPushButton('‹')
        previous.setFixedWidth(30)
        previous.clicked.connect(lambda: self.shift_date(-1))
        following=self.next_day=QPushButton('›')
        following.setFixedWidth(30)
        following.clicked.connect(lambda: self.shift_date(1))
        self.view=QComboBox()
        self.view.addItems(['Классы','Учителя','Кабинеты','Группы'])
        self.cell_width=QSpinBox()
        self.cell_width.setRange(28,100)
        self.cell_width.setValue(40)
        self.cell_height=QSpinBox()
        self.cell_height.setRange(22,100)
        self.cell_height.setValue(25)
        self.search=QLineEdit()
        self.search.setPlaceholderText('Поиск: предмет, учитель, класс…')
        for widget in (self.layout_mode,self.scope,previous,self.date_edit,following,self.weekday,self.view,
                       QLabel('Ячейка'),self.cell_width,QLabel('Строка'),self.cell_height,self.search):bar.addWidget(widget)
        body.addLayout(bar)
        summary=QHBoxLayout()
        self.week_summary=QLabel()
        self.week_summary.setObjectName('muted')
        summary.addWidget(self.week_summary,1)
        summary.addWidget(QLabel('Цвет карточки: учитель'))
        self.issue_button=QPushButton('Замечания')
        self.issue_button.clicked.connect(self.open_control)
        summary.addWidget(self.issue_button)
        body.addLayout(summary)
        self.stack=QStackedWidget()
        self.week_grid=DeskTable(self)
        self.week_grid.lesson_opened.connect(self.edit_week_lesson)
        self.week_grid.lesson_moved.connect(self.move_week_lesson)
        self.stack.addWidget(self.week_grid)
        self.grid=Grid(self)
        self.grid.moved.connect(self.move)
        self.grid.cellDoubleClicked.connect(self.edit_cell)
        self.stack.addWidget(self.grid)
        body.addWidget(self.stack,1)
        self.pending_label=QLabel()
        self.pending_label.setObjectName('muted')
        body.addWidget(self.pending_label)
        self.pending_list=QListWidget()
        self.pending_list.setFlow(QListWidget.Flow.LeftToRight)
        self.pending_list.setWrapping(False)
        self.pending_list.setMaximumHeight(68)
        self.pending_list.itemDoubleClicked.connect(self.open_pending)
        body.addWidget(self.pending_list)
        self.info=QLabel()
        self.info.setObjectName('muted')
        body.addWidget(self.info)
        self.issues=QListWidget(self)
        self.issues.hide()
        layout.addWidget(content,1)
        self.setCentralWidget(central)
        self.setStyleSheet(STYLE)
        menus = {
            'Проект': [('Демонстрационный пример', self.demo), ('Экспорт проекта…', self.export),
                        ('Импорт проекта…', self.import_archive), ('Резервная копия…', self.backup),
                        ('Импорт XLSX…', self.import_excel), ('Экспорт XLSX…', self.excel),
                        ('Загрузить школьную книгу и кабинеты…', self.import_school), ('Открыть сохранённую сверку…', self.open_school),
                        ('Загрузить сохранённую книгу в расписание…', self.activate_school)],
            'Генерация': [('Разместить занятия базовой недели…', lambda: self.calculate(weekly=True))],
            'Печать': [('Предварительный просмотр…', self.print_preview), ('Сохранить PDF…', self.pdf)],
            'Доступ': [('Создать пользователя…', self.user_create), ('Сменить пользователя', self.logout)],
        }
        for title, entries in menus.items():
            menu = self.menuBar().addMenu(title)
            for name, callback in entries:
                menu.addAction(name, callback)
        for key, callback in [(QKeySequence.StandardKey.Undo, self.undo), (QKeySequence.StandardKey.Redo, self.redo)]:
            action = QAction(self)
            action.setShortcut(QKeySequence(key))
            action.triggered.connect(callback)
            self.addAction(action)
        for widget in (self.view, self.scope, self.weekday, self.layout_mode):
            widget.currentIndexChanged.connect(self.refresh)
        self.date_edit.dateChanged.connect(self.refresh)
        self.search.textChanged.connect(self.refresh)
        self.cell_width.valueChanged.connect(self.refresh)
        self.cell_height.valueChanged.connect(self.refresh)
        self.refresh()

    def open_pending(self,item):
        self.scope.setCurrentIndex(1)
        key=item.data(Qt.ItemDataRole.UserRole)
        lesson=next(e for e in self.p['lessons'] if e['id']==key)
        self.edit_week_lesson(key,None,lesson['weekday'])

    def extract_teacher_assignments(self):
        def work():
            self.store.require('catalog')
            version,p=self.store.load()
            if not p.get('source_workbook'):
                raise RuleError('Сначала импортируйте школьную книгу')
            from .teacher_import import resolve_teachers
            from .school_import import stable_id
            source=resolve_teachers(p['source_workbook'])
            by_id={e['id']:e for e in source['events']}
            count=0
            updated=deepcopy(p)
            for lesson in updated['lessons']:
                event=by_id.get(lesson['id'])
                if lesson['teachers'] or not event or not event.get('teacher'):continue
                tid=stable_id('teacher',event['teacher'])
                if tid not in updated['teachers']:continue
                lesson['teachers']=[tid]
                lesson['teacher_source']=event['teacher_source']
                subjects=updated['teachers'][tid]['subjects']
                if lesson['subject'] not in subjects:subjects.append(lesson['subject'])
                count+=1
            if not count:
                QMessageBox.information(self,'Назначения','Новых однозначных назначений не найдено.')
                return
            if not confirm(self,f'Применить {count} назначений из книги к занятиям без преподавателя? Существующие назначения сохранятся. Будет создана резервная копия.'):return
            updated['source_workbook']=source
            self.automatic_backup()
            self.store.save(updated,version,'school.teachers','catalog')
            self.refresh()
        self.run(work)

    def open_control(self):
        self.show_text('Контроль расписания', '\n'.join(self.issues.item(i).text() for i in range(self.issues.count())))

    def open_replacements(self):
        self.scope.setCurrentIndex(0)
        self.layout_mode.setCurrentIndex(1)

    def shift_date(self, direction):
        self.date_edit.setDate(self.date_edit.date().addDays(direction*(7 if self.layout_mode.currentIndex()==0 else 1)))

    def edit_week_lesson(self, lesson_id, day, weekday):
        if day:
            self.date_edit.setDate(QDate.fromString(day,'yyyy-MM-dd'))
            lesson=next(e for e in day_lessons(self.p,day) if e['id']==lesson_id)
        else:
            self.weekday.setCurrentIndex(weekday-1)
            lesson=next(e for e in self.p['lessons'] if e['id']==lesson_id)
        self.lesson_form(lesson)

    def move_week_lesson(self, lesson_id, day, weekday, slot):
        def work():
            if day:
                if getattr(self.week_grid,'drag_day',day)!=day:
                    raise RuleError('Для переноса между днями выберите базовую неделю')
                if not any(e['id']==lesson_id for e in day_lessons(self.p,day)):
                    raise RuleError('Для переноса между днями выберите базовую неделю')
                self.date_edit.setDate(QDate.fromString(day,'yyyy-MM-dd'))
            self.weekday.setCurrentIndex(weekday-1)
            self.move(lesson_id,slot)
        self.run(work)

    def eventFilter(self, watched, event):
        if event.type() in (QEvent.Type.KeyPress, QEvent.Type.MouseButtonPress, QEvent.Type.MouseMove):
            self.last_activity = time.monotonic()
        return False

    def check_idle(self):
        if not self.locking and time.monotonic() - self.last_activity > 15 * 60:
            self.logout()

    def logout(self):
        self.locking = True
        self.cancel_worker()
        self.hide()
        self.store.user = None
        self.undo_stack.clear()
        self.redo_stack.clear()
        dialog = Login(self.store)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            self.close()
        else:
            self.last_activity = time.monotonic()
            self.refresh()
            self.show()
        self.locking = False

    def run(self, callback):
        try:
            return callback()
        except Exception as exc:
            alert(self, exc)

    def current_day(self):
        return self.date_edit.date().toString('yyyy-MM-dd') if self.scope.currentIndex() == 0 else None

    def refresh(self, *_):
        try:
            self.version, self.p = self.store.load()
        except RuleError:
            return
        day = self.current_day()
        weekday = self.weekday.currentIndex() + 1
        self.date_edit.setVisible(day is not None)
        self.previous_day.setEnabled(day is not None)
        self.next_day.setEnabled(day is not None)
        self.weekday.setVisible(day is None and self.layout_mode.currentIndex()==1)
        self.lessons = day_lessons(self.p, day) if day else [x for x in self.p['lessons'] if x['weekday'] == weekday]
        schedule = self.p['days'].get(day, {})
        status = STATUSES.get(schedule.get('status'), 'Из базового шаблона') if day else 'Базовая неделя · редактирование'
        self.banner.setText(('ДЕМОНСТРАЦИОННЫЕ ДАННЫЕ  •  ' if self.p['demo'] else '') + self.p['school'] + '  •  ' + self.p['year'])
        self.banner.setStyleSheet('color:#875519; padding:10px; background:#FFF3D8; border-radius:6px' if self.p['demo'] else 'color:#53657D')
        self.info.setText(f"{status} · редакция {self.version} · {self.store.user['login']} · {len(self.p['teachers'])} учителей · изменения сохранены")
        view = self.view.currentText()
        kind = {'Учителя': 'teachers', 'Кабинеты': 'rooms', 'Группы': 'groups'}.get(view)
        columns = [(k, v['name']) for k, v in self.p[kind].items()] if kind else [(c, c) for c in sorted({v['class'] for v in self.p['groups'].values()}, key=class_sort)]
        self.grid.clear()
        self.grid.setRowCount(self.p['periods'])
        self.grid.setColumnCount(len(columns))
        self.grid.setHorizontalHeaderLabels([v for k, v in columns])
        self.grid.setVerticalHeaderLabels([f'{i} урок' for i in range(1, self.p['periods'] + 1)])
        base = {x['id']: x for x in schedule.get('source', [])}
        search = self.search.text().casefold()
        for row in range(self.p['periods']):
            self.grid.setRowHeight(row, 108)
            for col, (key, name) in enumerate(columns):
                found = []
                for lesson in self.lessons:
                    match = (key in lesson[kind] if kind in ('teachers', 'groups') else lesson['room'] == key if kind == 'rooms' else any(self.p['groups'][g]['class'] == key for g in lesson['groups']))
                    text = self.lesson_text(lesson)
                    if match and lesson['slot'] <= row + 1 < lesson['slot'] + lesson.get('duration', 1) and (not search or search in text.casefold()):
                        found.append(lesson)
                item = QTableWidgetItem('\n────────────\n'.join(self.lesson_text(x) for x in found))
                item.setData(Qt.ItemDataRole.UserRole, [x['id'] for x in found])
                item.setToolTip(item.text() or 'Свободный слот')
                item.setData(CARDS_ROLE, self.daily_cards(found, base, day))
                if found:
                    color = found[0].get('color', '#E2EDFA') if not found[0]['teachers'] else self.p['teachers'][found[0]['teachers'][0]].get('color', '#E2EDFA')
                    item.setBackground(QColor(color))
                    item.setForeground(QColor('#142D4E'))
                    if day and any(x['id'] not in base or x != base[x['id']] for x in found) and base:
                        item.setText('Δ ИЗМЕНЕНО\n' + item.text())
                    if any(x.get('status') == 'cancelled' for x in found):
                        item.setBackground(QColor('#FDE1E1'))
                self.grid.setItem(row, col, item)
                if len(found) >= 3:
                    self.grid.setRowHeight(row, 126)
        errors = validate(self.p, self.lessons if day else self.p['lessons'], day)
        self.issues.clear()
        self.issues.addItems(errors or ['✓ Пересечений и ресурсных конфликтов нет'])
        if not self.p['rules']['approved']:
            self.issues.addItem('! Нормативный профиль школы не утверждён')
        self.stack.setCurrentIndex(self.layout_mode.currentIndex())
        self.week_grid.populate(self.p,self.view.currentText(),day,self.search.text(),
                                self.cell_width.value(),self.cell_height.value(),self.lesson_text)
        self.week_summary.setText(f"{len(self.p['lessons'])} занятий · {len({g['class'] for g in self.p['groups'].values()})} классов · {len(self.p['teachers'])} учителей")
        self.issue_button.setText(f'{len(errors)} замечаний')
        pending=[e for e in self.p['lessons'] if not e['teachers']]
        self.pending_label.setText(f'Требуют назначения преподавателя · {len(pending)} занятий · двойной щелчок — открыть')
        self.pending_list.clear()
        for lesson in pending:
            groups=', '.join(sorted({self.p['groups'][g]['class'] for g in lesson['groups']},key=class_sort))
            item=QListWidgetItem(f"{groups} · {DAYS[lesson['weekday']-1][:2]} · {lesson['slot']}\n{self.p['subjects'][lesson['subject']]['name']}")
            item.setSizeHint(QSize(145,48))
            item.setData(Qt.ItemDataRole.UserRole,lesson['id'])
            item.setToolTip(self.lesson_text(lesson))
            item.setBackground(QColor(lesson.get('color','#FFF0D2')))
            self.pending_list.addItem(item)
        self.pending_list.setVisible(bool(pending))
        self.pending_label.setVisible(bool(pending))
        self.statusBar().showMessage('Двойной щелчок — редактировать · перетаскивание — изменить время · Ctrl+Z / Ctrl+Y — отмена / повтор')

    def daily_cards(self, lessons, base, day):
        cards = []
        for lesson in lessons:
            status = 'ОТМЕНА · ' if lesson.get('status') == 'cancelled' else ''
            changed = day and base and (lesson['id'] not in base or lesson != base[lesson['id']])
            label = ('Δ ' if changed else '') + status + self.p['subjects'][lesson['subject']]['name']
            label += '\n' + (', '.join(self.p['teachers'][t]['name'] for t in lesson['teachers']) or 'Преподаватель не назначен')
            label += ' · ' + self.p['rooms'].get(lesson['room'], {}).get('name', 'Кабинет не назначен')
            color = '#FDE1E1' if status else (self.p['teachers'][lesson['teachers'][0]].get('color', '#E2EDFA') if lesson['teachers'] else lesson.get('color', '#FFF3D8'))
            cards.append(dict(id=lesson['id'], label=label, color=color, tooltip=self.lesson_text(lesson)))
        return cards

    def lesson_text(self, l):
        marker = 'ОТМЕНА · ' if l.get('status') == 'cancelled' else 'ВОСПОЛНЕНИЕ · ' if l.get('status') == 'makeup' else ''
        return marker + self.p['subjects'][l['subject']]['name'] + (' 🔒' if l.get('locked') else '') + '\n' + ', '.join(self.p['groups'][g]['name'] for g in l['groups']) + '\n' + (', '.join(self.p['teachers'][t]['name'] for t in l['teachers']) or 'Преподаватель не назначен') + ' · каб. ' + self.p['rooms'].get(l['room'], {}).get('name', 'не назначен')

    def checkpoint(self):
        self.undo_stack.append((self.version, deepcopy(self.p)))
        self.redo_stack.clear()

    def move(self, lesson_id, slot):
        def work():
            previous = deepcopy(self.p)
            self.service.move(self.current_day(), lesson_id, slot, self.weekday.currentIndex() + 1)
            self.undo_stack.append((self.store.load()[0], previous))
            self.redo_stack.clear()
            self.refresh()
        self.run(work)

    def preview_move(self, lesson_id):
        day = self.current_day()
        original = day_lessons(self.p, day) if day else deepcopy(self.p['lessons'])
        for row in range(self.grid.rowCount()):
            lessons = deepcopy(original)
            lesson = next(x for x in lessons if x['id'] == lesson_id)
            lesson['slot'] = row + 1
            errors = ['Урок закреплён'] if lesson.get('locked') else validate(self.p, lessons, day, draft=True)
            for col in range(self.grid.columnCount()):
                item = self.grid.item(row, col)
                item.setData(PREVIEW_ROLE, not bool(errors))
                item.setBackground(QColor('#FDE1E1' if errors else '#D2F0E5'))
                item.setToolTip('\n'.join(errors) if errors else '✓ Перенос допустим')

    def undo(self):
        self.restore_stack(self.undo_stack, self.redo_stack, 'edit.undo')

    def redo(self):
        self.restore_stack(self.redo_stack, self.undo_stack, 'edit.redo')

    def restore_stack(self, source, target, action):
        def work():
            self.store.require('edit')
            if not source:
                return
            current, p = self.store.load()
            expected, old = source[-1]
            if current != expected:
                raise RuleError('После этой операции данные изменились. Откат заблокирован для сохранности новых записей.')
            # Undo only editor changes; never revert approvals, accounts or obligations.
            if p['obligations'] != old['obligations'] or p['outbox'] != old['outbox']:
                raise RuleError('Изменились обязательства или публикации; используйте отдельную корректировку дня')
            version = self.store.save(old, current, action)
            source.pop()
            target.append((version, p))
            if source:
                source[-1] = (version, source[-1][1])
            self.refresh()
        self.run(work)

    def edit_cell(self, row, column):
        ids = self.grid.item(row, column).data(Qt.ItemDataRole.UserRole)
        if not ids:
            self.add_lesson()
            return
        if self.grid.hit_lesson in ids:
            key = self.grid.hit_lesson
        elif len(ids) > 1:
            labels = [self.lesson_text(next(x for x in self.lessons if x['id'] == key)).replace('\n', ' / ') for key in ids]
            selected, ok = QInputDialog.getItem(self, 'Выберите карточку', 'Занятие', labels, 0, False)
            if not ok:
                return
            key = ids[labels.index(selected)]
        else:
            key = ids[0]
        self.lesson_form(next(x for x in self.lessons if x['id'] == key))

    def add_lesson(self):
        self.lesson_form()

    def lesson_form(self, existing=None):
        def work():
            self.store.require('edit')
            if not all(self.p[k] for k in ('subjects', 'groups')):
                raise RuleError('Сначала заполните справочники или загрузите демонстрационный пример')
            d = Form(self, 'Занятие' if existing else 'Новое занятие')
            subject = d.combo('Предмет', [(k, v['name']) for k, v in self.p['subjects'].items()])
            groups = d.text('Группы: ID через запятую', ','.join(existing['groups']) if existing else next(iter(self.p['groups'])))
            teachers = QListWidget()
            teachers.setMaximumHeight(160)
            for tid, record in self.p['teachers'].items():
                item = QListWidgetItem(record['name'])
                item.setData(Qt.ItemDataRole.UserRole, tid)
                item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
                item.setCheckState(Qt.CheckState.Checked if existing and tid in existing['teachers'] else Qt.CheckState.Unchecked)
                teachers.addItem(item)
            d.form.addRow('Преподаватели', teachers)
            qualification = d.check('Подтверждаю допуск выбранных преподавателей к этому предмету')
            if existing and existing.get('teacher_source'):
                evidence=QLabel('Назначение: '+existing['teacher_source'])
                evidence.setWordWrap(True)
                d.form.addRow(evidence)
            if existing and existing.get('teacher_candidates'):
                d.form.addRow(QLabel('Варианты из исходного снимка: ' + ', '.join(existing['teacher_candidates'])))
            room = d.combo('Кабинет', [(None, 'Не назначен')] + [(k, v['name']) for k, v in self.p['rooms'].items()])
            slot = d.number('Номер урока', existing['slot'] if existing else 1, self.p['periods'])
            duration = d.number('Длительность в слотах', existing.get('duration', 1) if existing else 1, self.p['periods'])
            locked = d.check('Закрепить от перемещения', existing.get('locked', False) if existing else False)
            cancelled = d.check('Отменено (укажите основание)', existing.get('status') == 'cancelled' if existing else False)
            comment = d.text('Комментарий / основание', existing.get('comment', '') if existing else '')
            for widget, field in ((subject, 'subject'), (room, 'room')):
                if existing:
                    widget.setCurrentIndex(widget.findData(existing[field]))
            d.form.addRow(QLabel('ID справочников доступны в «Справочники и правила».\nОбъединённые группы и совместные учителя — одна запись.'))
            if not d.finish():
                return
            version, p = self.store.load()
            if version != self.version:
                raise RuleError('Данные изменились. Обновите экран')
            lesson = dict(id=existing['id'] if existing else uid(), weekday=self.date_edit.date().dayOfWeek() if self.current_day() else self.weekday.currentIndex() + 1,
                          slot=slot.value(), duration=duration.value(), subject=subject.currentData(),
                          groups=[s.strip() for s in groups.text().split(',')], teachers=[teachers.item(i).data(Qt.ItemDataRole.UserRole) for i in range(teachers.count()) if teachers.item(i).checkState() == Qt.CheckState.Checked],
                          room=room.currentData(), locked=locked.isChecked(), status='cancelled' if cancelled.isChecked() else 'planned', comment=comment.text())
            if qualification.isChecked():
                for tid in lesson['teachers']:
                    if lesson['subject'] not in p['teachers'][tid]['subjects']:
                        p['teachers'][tid]['subjects'].append(lesson['subject'])
            if existing:
                lesson = {**existing, **lesson}
            if cancelled.isChecked() and not comment.text().strip():
                raise RuleError('Укажите основание отмены')
            day = self.current_day()
            lessons = day_lessons(p, day) if day else deepcopy(p['lessons'])
            lessons = [l for l in lessons if l['id'] != lesson['id']] + [lesson]
            if day:
                self.service.edit_day(p, day, lessons)
            else:
                p['lessons'] = lessons
            result = self.store.save(p, version, 'lesson.edit')
            self.undo_stack.append((result, deepcopy(self.p)))
            self.redo_stack.clear()
            self.refresh()
        self.run(work)

    def demo(self):
        def work():
            self.store.require('catalog')
            if not confirm(self, 'Заменить текущий проект демонстрационным? Сначала будет создана резервная копия.'):
                return
            self.automatic_backup()
            version, _ = self.store.load()
            self.store.save(demo_project(), version, 'project.demo', 'catalog')
            self.refresh()
        self.run(work)

    def catalog(self):
        def work():
            self.store.require('catalog')
            dialog = CatalogDialog(self, self.p)
            if dialog.exec() != QDialog.DialogCode.Accepted:
                return
            version, p = self.store.load()
            if version != self.version:
                raise RuleError('Проект изменился, откройте справочники заново')
            changed = dialog.result_project
            if changed['rules'] != p['rules']:
                self.store.require('approve')
            self.store.save(changed, version, 'catalog.edit', 'catalog')
            self.refresh()
        self.run(work)

    def absence(self):
        def work():
            self.store.require('edit')
            d = Form(self, 'Регистрация отсутствия')
            teacher = d.combo('Преподаватель', [(k, v['name']) for k, v in self.p['teachers'].items()])
            start, end = QDateEdit(self.date_edit.date()), QDateEdit(self.date_edit.date())
            for label, widget in [('С', start), ('По', end)]:
                widget.setCalendarPopup(True)
                d.form.addRow(label, widget)
            slots = d.text('Номера уроков (пусто = весь день)')
            reason = d.combo('Служебная причина', [(s, s) for s in ['Больничный', 'Отпуск', 'Отгул', 'Служебное отсутствие']])
            d.form.addRow(QLabel('Медицинские подробности не записываются.'))
            if d.finish():
                periods = [int(s.strip()) for s in slots.text().split(',') if s.strip()]
                self.service.add_absence(teacher.currentData(), start.date().toString('yyyy-MM-dd'), end.date().toString('yyyy-MM-dd'), periods, reason.currentData())
                self.refresh()
        self.run(work)

    def calculate(self, _checked=False, weekly=False):
        def work():
            self.store.require('solve')
            if not weekly and not self.current_day():
                raise RuleError('Для замен выберите режим «На дату»')
            if self.worker:
                raise RuleError('Расчёт уже выполняется')
            self.solve_version, project = self.store.load()
            self.solving_week = weekly
            pending = validate(project, project['lessons'])
            if any('не назначен' in e or 'не задана' in e or 'требует подтверждения' in e for e in pending):
                raise RuleError('Перед расчётом заполните назначения, вместимость кабинетов и составы групп. Список — в контроле расписания.')
            if weekly and not project['lessons']:
                raise RuleError('Введите занятия базовой недели: одна запись на каждое недельное занятие. Генератор разместит их по доступным дням и слотам.')
            self.solve_day = self.current_day()
            self.store.log('solver.start', self.solve_day)
            context = mp.get_context('spawn')
            self.connection, child = context.Pipe(duplex=False)
            self.worker = context.Process(target=_solve_worker, args=(child, project, self.solve_day, weekly), daemon=True)
            self.worker.start()
            child.close()
            self.progress = QProgressDialog('Размещение занятий недели…' if weekly else 'Рассчитываются два независимых варианта…', 'Отменить', 0, 0, self)
            self.progress.setWindowModality(Qt.WindowModality.WindowModal)
            self.progress.canceled.connect(self.cancel_worker)
            self.progress.show()
            self.poll = QTimer(self)
            self.poll.timeout.connect(self.poll_worker)
            self.poll.start(150)
        self.run(work)

    def cancel_worker(self):
        if self.worker:
            self.poll.stop()
            self.worker.terminate()
            self.worker.join(timeout=2)
            self.worker = None
            self.connection.close()
            self.progress.close()

    def poll_worker(self):
        if self.connection.poll():
            try:
                status, result = self.connection.recv()
            except EOFError:
                status, result = 'error', 'Расчёт завершился без результата'
            self.poll.stop()
            self.worker.join(timeout=2)
            self.worker = None
            self.connection.close()
            self.progress.close()
            if status == 'error':
                alert(self, result)
            elif self.solving_week:
                def apply_week():
                    if not result['feasible']:
                        raise RuleError(result['diagnostic'])
                    if not confirm(self, f"Найдено допустимое расписание: {len(result['lessons'])} занятий. Применить к базовой неделе? Дневные снимки сохранятся."):
                        return
                    version, p = self.store.load()
                    if version != self.solve_version:
                        raise RuleError('Расчёт устарел, повторите генерацию')
                    p['lessons'] = result['lessons']
                    self.store.save(p, version, 'week.generate')
                    self.refresh()
                self.run(apply_week)
            else:
                self.show_scenarios(result)
        elif not self.worker.is_alive():
            self.cancel_worker()
            alert(self, 'Процесс расчёта завершился без результата')

    def show_scenarios(self, results):
        dialog = Form(self, 'Сравнение вариантов · исходная редакция ' + str(self.solve_version))
        dialog.resize(800, 650)
        choice = dialog.combo('Выбрать', [('A', 'A — бережный: сохраняет номера уроков'), ('B', 'B — расширенный: перестраивает день')])
        text = QTextEdit()
        text.setReadOnly(True)
        parts = []
        for result in results:
            parts.append(f"РЕЖИМ {result['mode']} · {result.get('status', '')}\n{result['diagnostic']}")
            parts.extend(f'{k}: {v}' for k, v in result['metrics'].items())
            originals = {x['id']: x for x in day_lessons(self.p, self.solve_day)}
            for lesson in result['lessons']:
                before = originals[lesson['id']]
                if before != lesson:
                    parts.append('ДО: ' + self.lesson_text(before).replace('\n', ' / ') + f" · урок {before['slot']}")
                    parts.append('ПОСЛЕ: ' + self.lesson_text(lesson).replace('\n', ' / ') + f" · урок {lesson['slot']}")
            parts.append('\n')
        text.setPlainText('\n'.join(parts))
        dialog.form.addRow(text)
        makeup = dialog.check('Создать обязательства для непроведённых уроков (по правилу школы)', False)
        dialog.form.addRow(QLabel('Применение сохраняет черновик. Утверждение — отдельное действие.'))
        if dialog.finish():
            def work():
                selected = next(x for x in results if x['mode'] == choice.currentData())
                self.service.apply_scenario(self.solve_day, selected, self.solve_version, makeup.isChecked())
                self.refresh()
            self.run(work)

    def transition(self, target):
        def work():
            if not self.current_day():
                raise RuleError('Выберите конкретную дату')
            self.service.transition(self.current_day(), target)
            self.refresh()
            if target == 'published':
                QMessageBox.information(self, 'Публикация', 'Локальный выпуск сформирован. Внешний канал не настроен, отправка не выполнялась. Печать и PDF доступны.')
        self.run(work)

    def obligations(self):
        def work():
            self.refresh()
            opened = [o for o in self.p['obligations'] if o['status'] != 'closed']
            if not opened:
                QMessageBox.information(self, 'Восполнения', 'Открытых обязательств нет.')
                return
            labels = [o['source_day'] + ' / ' + o['status'] + ' / ' + self.lesson_text(o['lesson']).replace('\n', ' / ') for o in opened]
            selection, ok = QInputDialog.getItem(self, 'Восполнения', 'Выберите обязательство', labels, 0, False)
            if not ok:
                return
            obligation = opened[labels.index(selection)]
            if obligation['status'] == 'assigned':
                if confirm(self, 'Подтверждаете фактическое проведение назначенного восполнения? Это закроет обязательство.'):
                    self.service.close_makeup(obligation['id'])
            else:
                d = Form(self, 'Назначить восполнение')
                target = QDateEdit(self.date_edit.date().addDays(1))
                target.setCalendarPopup(True)
                d.form.addRow('Дата', target)
                slot = d.number('Урок', 1, self.p['periods'])
                teacher = d.combo('Учитель', [(k, v['name']) for k, v in self.p['teachers'].items()])
                room = d.combo('Кабинет', [(None, 'Не назначен')] + [(k, v['name']) for k, v in self.p['rooms'].items()])
                if d.finish():
                    self.service.assign_makeup(obligation['id'], target.date().toString('yyyy-MM-dd'), slot.value(), teacher.currentData(), room.currentData())
            self.refresh()
        self.run(work)

    def show_text(self, title, text):
        d = QDialog(self)
        d.setWindowTitle(title)
        d.resize(900, 600)
        layout = QVBoxLayout(d)
        edit = QTextEdit()
        edit.setReadOnly(True)
        edit.setPlainText(text)
        layout.addWidget(edit)
        d.exec()

    def history(self):
        self.run(lambda: self.show_text('Журнал действий', '\n'.join(' | '.join(str(v or '') for v in row) for row in self.store.history())))

    def outbox(self):
        self.show_text('Локальная очередь · внешние адаптеры ещё не подключены', '\n'.join(f"{i['day']} · R{i['revision']} · {i['status']} · {i['channel']}" for i in self.p['outbox']) or 'Очередь пуста')

    def automatic_backup(self):
        self.store.backup(self.store.path.parent / ('backup-' + datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '.db'))

    def backup(self):
        path, _ = QFileDialog.getSaveFileName(self, 'Резервная копия', 'TimeLabs-backup.db', 'SQLite (*.db)')
        if path:
            self.run(lambda: self.store.backup(path))

    def export(self):
        path, _ = QFileDialog.getSaveFileName(self, 'Экспорт проекта без пользователей и паролей', 'TimeLabs.tlabs', 'TimeLabs (*.tlabs)')
        if path:
            self.run(lambda: self.store.export_project(path))

    def import_archive(self):
        path, _ = QFileDialog.getOpenFileName(self, 'Импорт проекта', '', 'TimeLabs (*.tlabs)')
        if path and confirm(self, 'Заменить текущий проект? Будет создана резервная копия. Импортированные дни требуют повторного утверждения.'):
            self.run(lambda: self.store.import_project(path))
            self.refresh()

    def excel(self):
        path, _ = QFileDialog.getSaveFileName(self, 'Экспорт базовой недели', 'Расписание.xlsx', 'Excel (*.xlsx)')
        if path:
            self.run(lambda: export_xlsx(self.p, self.p['lessons'], path))

    def import_school(self, _checked=False, workbook_path=None):
        def work():
            self.store.require('catalog')
            path = workbook_path
            if not path:
                path, _ = QFileDialog.getOpenFileName(self, 'Уточнённая школьная книга', '', 'Excel (*.xlsx)')
            if not path:
                return
            room_path, _ = QFileDialog.getOpenFileName(self, 'Список кабинетов (можно пропустить)', '', 'Word (*.docx)')
            source = read_school_workbook(path, room_path or None)
            _, current = self.store.load()
            previous = current.get('source_workbook', {})
            if previous.get('sha256') == source['sha256']:
                for clarification in previous.get('clarifications', []):
                    source = apply_clarification(source, clarification)
            dialog = SourceDialog(source, self, can_save=True, save_label='Загрузить в рабочее расписание')
            if dialog.exec() == QDialog.DialogCode.Accepted:
                version, project = self.store.load()
                project = self.prepare_school_transfer(project, dialog.source)
                if project is None:
                    return
                self.store.save(project, version, 'school.import', 'catalog')
                self.refresh()
        self.run(work)

    def prepare_school_transfer(self, project, source):
        if not confirm(self, f"Загрузить {len(source['events'])} занятий в рабочее расписание? Текущее расписание и справочники будут заменены. Будет создана резервная копия. Незаданные назначения останутся доступными для редактирования."):
            return None
        result = transfer_school(project, source)
        self.automatic_backup()
        return result

    def activate_school(self):
        def work():
            self.store.require('catalog')
            version, project = self.store.load()
            if not project.get('source_workbook'):
                raise RuleError('Сначала загрузите школьную книгу')
            result = self.prepare_school_transfer(project, project['source_workbook'])
            if result is not None:
                self.store.save(result, version, 'school.activate', 'catalog')
                self.scope.setCurrentIndex(1)
                self.refresh()
        self.run(work)

    def open_school(self):
        def work():
            self.store.require('view')
            expected, project = self.store.load()
            if not project.get('source_workbook'):
                raise RuleError('Сначала откройте школьную книгу через меню «Проект»')
            editable = 'catalog' in json.loads(self.store.user['permissions'])
            dialog = SourceDialog(project['source_workbook'], self, can_save=editable)
            if dialog.exec() == QDialog.DialogCode.Accepted:
                version, project = self.store.load()
                if version != expected:
                    raise RuleError('Проект изменился во время сверки. Откройте его повторно')
                project['source_workbook'] = dialog.source
                self.store.save(project, expected, 'school.clarification', 'catalog')
                self.refresh()
        self.run(work)

    def import_excel(self):
        def work():
            self.store.require('catalog')
            path, _ = QFileDialog.getOpenFileName(self, 'Импорт расписания', '', 'Расписание (*.xlsx *.tlabs);;Excel (*.xlsx);;TimeLabs (*.tlabs)')
            if not path:
                return
            if Path(path).suffix.lower()=='.tlabs':
                if confirm(self,'Заменить текущий проект? Перед импортом будет создана резервная копия.'):
                    self.store.import_project(path)
                    self.refresh()
                return
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True)
            try:
                is_school = 'Группы — записи' in book.sheetnames
            finally:
                book.close()
            if is_school:
                self.import_school(workbook_path=path)
                return
            lessons, issues = preview_xlsx(self.p, path)
            self.show_text('Предварительная проверка импорта', f'Прочитано занятий: {len(lessons)}\n' + '\n'.join(issues or ['Ошибок не обнаружено. Изменения ещё не сохранены.']))
            if issues or not lessons or not confirm(self, f'Заменить базовую неделю: {len(lessons)} занятий? Существующие дневные снимки сохранятся.'):
                return
            version, p = self.store.load()
            if version != self.version:
                raise RuleError('Проект изменился во время импорта')
            self.automatic_backup()
            p['lessons'] = lessons
            self.store.save(p, version, 'xlsx.import', 'catalog')
            self.refresh()
        self.run(work)

    def document(self):
        day = self.current_day()
        if not day or self.p['days'].get(day, {}).get('status') not in ('approved', 'published'):
            raise RuleError('Для выдачи расписания сначала проверьте и утвердите конкретный день')
        doc = QTextDocument()
        doc.setDefaultFont(QFont('DejaVu Sans', 10))
        doc.setHtml(schedule_html(self.p, self.lessons, self.p['school'] + ' · ' + day,
                                 f"Редакция {self.p['days'][day]['revision']} · {STATUSES[self.p['days'][day]['status']]}"))
        return doc

    def printer(self):
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setPageOrientation(QPageLayout.Orientation.Landscape)
        return printer

    def pdf(self):
        def work():
            doc = self.document()
            path, _ = QFileDialog.getSaveFileName(self, 'Сохранить PDF', 'Расписание.pdf', 'PDF (*.pdf)')
            if path:
                printer = self.printer()
                printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
                printer.setOutputFileName(path)
                doc.print_(printer)
                self.store.log('print.pdf')
        self.run(work)

    def print_preview(self):
        def work():
            doc = self.document()
            preview = QPrintPreviewDialog(self.printer(), self)
            preview.paintRequested.connect(doc.print_)
            preview.exec()
        self.run(work)

    def user_create(self):
        def work():
            self.store.require('users')
            d = Form(self, 'Новый локальный пользователь')
            login = d.text('Логин')
            password = d.text('Пароль', password=True)
            role = d.combo('Роль', list(ROLE_NAMES.items()))
            if d.finish():
                self.store.create_user(login.text(), password.text(), role.currentData())
        self.run(work)

    def closeEvent(self, event):
        self.cancel_worker()
        event.accept()


def data_directory():
    if os.environ.get('TIMELABS_DATA_DIR'):
        return Path(os.environ['TIMELABS_DATA_DIR'])
    if sys.platform == 'win32':
        return Path(os.environ.get('LOCALAPPDATA', str(Path.home()))) / 'TimeLabs'
    return Path.home() / '.local' / 'share' / 'TimeLabs'


def main():
    mp.freeze_support()
    app = QApplication(sys.argv)
    app.setApplicationName('TimeLabs')
    app.setFont(QFont('Segoe UI' if sys.platform == 'win32' else 'DejaVu Sans', 10))
    app.setStyleSheet('QMainWindow,QDialog{background:#F5F7FB} QPushButton{padding:8px 12px;border:1px solid #CAD4E3;border-radius:5px;background:white;color:#142D4E} QPushButton:hover{background:#E4EDFD} QLineEdit,QComboBox,QDateEdit,QSpinBox{padding:7px;background:white;color:#142D4E;border:1px solid #CAD4E3;border-radius:4px} QTableWidget{background:white;gridline-color:#DFE6F0;color:#142D4E} QHeaderView::section{padding:10px;background:#EAF0F8;color:#142D4E;border:0} QLabel{color:#142D4E}')
    directory = data_directory()
    directory.mkdir(parents=True, exist_ok=True)
    lock = QLockFile(str(directory / 'TimeLabs.lock'))
    lock.setStaleLockTime(0)
    if not lock.tryLock(100):
        alert(None, 'Эта база уже открыта в другой копии TimeLabs. Завершите её работу перед повторным запуском.')
        return 1
    try:
        store = Store(directory / 'timelabs.db')
        if Login(store).exec() != QDialog.DialogCode.Accepted:
            return 0
        window = Window(store)
        window.show()
        return app.exec()
    finally:
        lock.unlock()


if __name__ == '__main__':
    sys.exit(main())
